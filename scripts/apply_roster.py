"""回收《团队角色清单》Excel，写入 collab/roster.yaml 并重新生成协作文档。

用法：
    python scripts/apply_roster.py --xlsx docs/团队角色清单.xlsx

做三件事：
    1. 校验：角色代号必须在 roles.yaml 里，GitHub 用户名不能重复
    2. 写 collab/roster.yaml（人员信息单独存，不动 roles.yaml 的注释）
    3. 调 gen_collab_docs 重新生成《分工文件清单》与 CODEOWNERS
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

ROLES_FILE = ROOT / "collab" / "roles.yaml"
ROSTER_FILE = ROOT / "collab" / "roster.yaml"

HEADER_ROW = 4
COLS = {
    "code": 1,
    "level": 2,
    "title": 3,
    "blurb": 4,
    "files": 5,
    "branch": 6,
    "name": 7,
    "student_id": 8,
    "github": 9,
    "backup": 10,
    "skill_python": 11,
    "skill_sql": 12,
    "skill_web": 13,
    "skill_doc": 14,
    "hours": 15,
    "os": 16,
    "note": 17,
}


def cell(ws, row: int, key: str):
    v = ws.cell(row=row, column=COLS[key]).value
    if isinstance(v, str):
        v = v.strip()
    return v or None


def main() -> int:
    ap = argparse.ArgumentParser(description="回收角色清单并生成协作配置")
    ap.add_argument("--xlsx", default=str(ROOT / "docs" / "团队角色清单.xlsx"))
    args = ap.parse_args()

    path = Path(args.xlsx)
    if not path.exists():
        raise SystemExit(f"找不到文件：{path}")

    spec = yaml.safe_load(ROLES_FILE.read_text(encoding="utf-8"))
    valid = [r["code"] for r in spec["roles"]]

    wb = load_workbook(path, data_only=True)
    ws = wb["角色清单"]

    members: dict[str, dict] = {}
    problems: list[str] = []
    seen_github: dict[str, str] = {}

    for row in range(HEADER_ROW + 1, HEADER_ROW + 1 + len(valid) + 6):
        code = cell(ws, row, "code")
        if not code:
            continue
        if code not in valid:
            problems.append(f"第 {row} 行角色代号 `{code}` 不在 roles.yaml 里，已跳过")
            continue
        name = cell(ws, row, "name")
        github = cell(ws, row, "github")
        if not name:
            problems.append(f"{code}：姓名为空")
        if not github:
            problems.append(f"{code}：GitHub 用户名为空")
        else:
            key = github.lower().lstrip("@")
            if key in seen_github:
                problems.append(f"GitHub 用户名 `{github}` 与 {seen_github[key]} 重复")
            seen_github[key] = code
            github = "@" + key
        backup = cell(ws, row, "backup")
        if backup and backup not in valid:
            problems.append(f"{code}：第二志愿 `{backup}` 不是合法角色代号")
        members[code] = {
            "name": name,
            "student_id": cell(ws, row, "student_id"),
            "github": github,
            "backup": backup,
            "hours_per_week": cell(ws, row, "hours"),
            "os": cell(ws, row, "os"),
            "skill": {
                "python": cell(ws, row, "skill_python"),
                "sql_orm": cell(ws, row, "skill_sql"),
                "web": cell(ws, row, "skill_web"),
                "doc": cell(ws, row, "skill_doc"),
            },
            "note": cell(ws, row, "note"),
        }

    missing = [c for c in valid if c not in members]
    if missing:
        problems.append("以下角色没有任何行：" + "、".join(missing))

    payload = {"members": members}
    ROSTER_FILE.write_text(
        "# 人员信息（由 scripts/apply_roster.py 从 docs/团队角色清单.xlsx 生成，改人请改 Excel 后重跑）\n"
        + yaml.safe_dump(payload, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )
    print(f"已写入 {ROSTER_FILE.relative_to(ROOT)}，共 {len(members)} 人")

    import gen_collab_docs

    gen_collab_docs.main()

    if problems:
        print("\n需要处理的问题：")
        for p in problems:
            print("  ! " + p)
        return 1
    print("\n全部字段完整，CODEOWNERS 已可按路径指派评审人。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
